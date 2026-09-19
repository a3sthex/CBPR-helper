#!/usr/bin/env python3
"""NC//NET · отрисовка превью листа .xlsx в PNG (для показа дизайна без открытия файла).

Использование:
  PYTHONPATH=<deps> python3 tools/render_sheet_preview.py <файл.xlsx> <копия.xlsx> "<Лист>" out.png "[1,41]" "[1,9]" [масштаб]

<копия.xlsx> нужна только для движка formulas (по ней считаются значения формул) —
загружаемый и пересчитываемый файл должны быть одинаковыми. Требуется Pillow + formulas.
"""
import re, sys, json
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from PIL import Image, ImageDraw, ImageFont
from formulas import ExcelModel

XLSX, MODEL, SHEET, OUT = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
ROWS = range(*json.loads(sys.argv[5])); COLS = range(*json.loads(sys.argv[6]))
SCALE = float(sys.argv[7]) if len(sys.argv) > 7 else 1.25

wb = load_workbook(XLSX)
sol = ExcelModel().loads(MODEL).finish().calculate()

def clean(v):
    v = getattr(v, 'value', v)
    while isinstance(v, (list, tuple)) or type(v).__name__ == 'Array':
        try: v = v[0]
        except Exception: break
    s = str(v)
    s = re.sub(r'np\.(float|int)\d*\(([^()]*)\)', r'\2', s)
    s = s.replace('[[','').replace(']]','').replace('[','').replace(']','')
    s = s.strip('"')
    if re.fullmatch(r'-?0(\.0+)?', s): return '0'
    m = re.fullmatch(r'(-?\d+)\.0+', s)
    return m.group(1) if m else s

def computed(sheet, cell):
    for k in (f"'[{MODEL.split('/')[-1]}]{sheet}'!{cell}", f"[{MODEL.split('/')[-1]}]{sheet}'!{cell}"):
        if k in sol: return sol[k]
    return None

F='/usr/share/fonts/truetype/dejavu/'; CACHE={}
def fnt(mono, size, bold):
    k=(mono,size,bold)
    if k not in CACHE:
        n='DejaVuSansMono' if mono else 'DejaVuSans'
        CACHE[k]=ImageFont.truetype(F+n+('-Bold' if bold else '')+'.ttf', size)
    return CACHE[k]

ws = wb[SHEET]
rows, cols = list(ROWS), list(COLS)
widths={c: round((ws.column_dimensions[get_column_letter(c)].width or 8.43)*7*SCALE)+5 for c in cols}
heights={r: round((ws.row_dimensions[r].height or 15)*1.3333*SCALE) for r in rows}
W=sum(widths.values())+2; H=sum(heights.values())+2
img=Image.new('RGB',(W,H),'#0C0C10'); d=ImageDraw.Draw(img)
merged={}; skip=set()
for m in ws.merged_cells.ranges:
    c1,r1,c2,r2=range_boundaries(str(m)); merged[(r1,c1)]=(r2,c2)
    for r in range(r1,r2+1):
        for c in range(c1,c2+1):
            if (r,c)!=(r1,c1): skip.add((r,c))
y0=1
for r in rows:
    x=1
    for c in cols:
        cell=ws.cell(row=r,column=c); w=widths[c]; h=heights[r]
        if (r,c) in merged:
            r2,c2=merged[(r,c)]
            w=sum(widths[cc] for cc in cols if c<=cc<=c2) or w
            h=sum(heights[rr] for rr in rows if r<=rr<=r2) or h
        if (r,c) not in skip:
            try:
                fl=cell.fill
                if fl is not None and fl.fill_type=='solid' and isinstance(fl.start_color.rgb,str):
                    fg='#'+fl.start_color.rgb[-6:]
                    if fg!='#000000': d.rectangle([x,y0,x+w-1,y0+h-1],fill=fg)
            except Exception: pass
            txt=cell.value
            if isinstance(txt,str) and txt.startswith('='):
                v=computed(SHEET, cell.coordinate)
                txt='' if v is None else clean(v)
            if txt not in (None,''):
                try: fc='#'+cell.font.color.rgb[-6:]
                except Exception: fc='#E8E6E3'
                mono=(cell.font.name=='Courier New'); bold=bool(cell.font.bold)
                size=int(cell.font.size or 9); size=7 if size<=8 else 9
                f=fnt(mono,size,bold)
                hal=cell.alignment.horizontal or 'left'; val=cell.alignment.vertical or 'center'
                lines=[]; maxw=max(10,w-4)
                for part in str(txt).split('\n'):
                    part=part or ' '
                    while len(part)>0:
                        k=len(part)
                        while k>1 and d.textlength(part[:k],font=f)>maxw: k-=1
                        lines.append(part[:k]); part=part[k:]
                lh=size+2; total=len(lines)*lh
                ty=y0+2 if (val=='top' or total>=h-2) else y0+(h-total)//2
                for ln in lines:
                    if ty+size>y0+h: break
                    lw=d.textlength(ln,font=f)
                    tx=x+2 if hal=='left' else (x+(w-lw)/2 if hal=='center' else x+w-lw-2)
                    d.text((tx,ty),ln,font=f,fill=fc); ty+=lh
        x+=w
    y0+=h
img.save(OUT); print(f'{OUT} · {img.size[0]}×{img.size[1]}')
